"""
reporting/executive_summary.py
Generate a ~2-page executive summary DOCX from a SaveGeo analysis payload.
Business-oriented, not a technical report.
"""

from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# ── Color palette ──────────────────────────────────────────────────────────────
_C_BLUE   = RGBColor(0x1B, 0x52, 0x8A)
_C_GREEN  = RGBColor(0x2E, 0x7D, 0x32)
_C_RED    = RGBColor(0xB7, 0x1C, 0x1C)
_C_DARK   = RGBColor(0x21, 0x21, 0x21)
_C_GRAY   = RGBColor(0x75, 0x75, 0x75)
_C_WHITE  = RGBColor(0xFF, 0xFF, 0xFF)

_HEX_BLUE  = "1B528A"
_HEX_LBLUE = "EEF2F7"   # light blue row
_HEX_WHITE = "FFFFFF"
_HEX_GRAY  = "F5F5F5"

_FONT = "Aptos"   # falls back to Calibri on older Word


# ── Low-level helpers ──────────────────────────────────────────────────────────

def _fmt(val: Any, decimals: int = 1, fallback: str = "N/A") -> str:
    try:
        return f"{float(val):,.{decimals}f}"
    except (TypeError, ValueError):
        return fallback


def _fmt_area(ha: Any) -> str:
    try:
        v = float(ha)
        return f"{v:,.0f} ha" if v >= 1000 else f"{v:,.1f} ha"
    except (TypeError, ValueError):
        return "N/A"


def _set_font(run, size_pt: float, bold: bool = False, italic: bool = False,
              color: RGBColor | None = None) -> None:
    run.font.name = _FONT
    run.font.size = Pt(size_pt)
    run.font.bold  = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color


def _para(doc: Document, text: str = "", size_pt: float = 9.5,
          bold: bool = False, color: RGBColor | None = None,
          align: int | None = None,
          space_before: float = 0, space_after: float = 3) -> Any:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    if align is not None:
        p.alignment = align
    if text:
        run = p.add_run(text)
        _set_font(run, size_pt, bold=bold, color=color)
    return p


def _heading(doc: Document, text: str, size_pt: float = 10,
             color: RGBColor = _C_BLUE, border_hex: str = _HEX_BLUE,
             space_before: float = 6) -> Any:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(2)
    run = p.add_run(text.upper())
    _set_font(run, size_pt, bold=True, color=color)
    # bottom border line
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"),   "single")
    bottom.set(qn("w:sz"),    "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), border_hex)
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def _bullet(doc: Document, text: str, size_pt: float = 9.5,
            color: RGBColor | None = None) -> Any:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(2)
    p.paragraph_format.left_indent       = Cm(0.4)
    p.paragraph_format.first_line_indent = Cm(-0.4)
    run = p.add_run(f"• {text}")
    _set_font(run, size_pt, color=color)
    return p


def _shade_cell(cell, hex_color: str) -> None:
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"),  "clear")
    tcPr.append(shd)


def _fill_row(row, values: list[str], sizes: list[float],
              bold: bool = False, color: RGBColor | None = None,
              bg_hex: str | None = None) -> None:
    for cell, val, sz in zip(row.cells, values, sizes):
        if bg_hex:
            _shade_cell(cell, bg_hex)
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after  = Pt(1)
        run = p.add_run(str(val))
        _set_font(run, sz, bold=bold, color=color)


# ── Defensive extractors ───────────────────────────────────────────────────────

def get_nested(data: Any, path: str, default: Any = None) -> Any:
    """Traverse dot-separated path in nested dict; return default on miss."""
    val = data
    for key in path.split("."):
        if not isinstance(val, dict) or key not in val:
            return default
        val = val[key]
    return val if val is not None else default


def extract_landcover_rows(analysis: dict) -> list[dict]:
    """Top-5 land cover classes sorted by area, handling several response shapes."""
    lc = analysis.get("landcover") or {}
    rows: list[dict] = []

    if isinstance(lc.get("class_areas"), list):
        for item in lc["class_areas"]:
            rows.append({
                "class":   item.get("class") or item.get("name") or "Unknown",
                "area_ha": float(item.get("area_ha") or item.get("area") or 0),
                "percent": float(item.get("percent") or item.get("percentage") or 0),
            })
    elif isinstance(lc.get("classes"), dict):
        for cls, d in lc["classes"].items():
            if isinstance(d, dict):
                rows.append({
                    "class":   cls,
                    "area_ha": float(d.get("area_ha") or d.get("area") or 0),
                    "percent": float(d.get("percentage") or d.get("percent") or 0),
                })
    elif isinstance(lc.get("area_stats"), list):
        for item in lc["area_stats"]:
            rows.append({
                "class":   item.get("class") or item.get("label") or "Unknown",
                "area_ha": float(item.get("area_ha") or 0),
                "percent": float(item.get("percent") or item.get("percentage") or 0),
            })

    rows.sort(key=lambda x: x["area_ha"], reverse=True)
    return rows[:5]


def extract_transition_rows(analysis: dict) -> list[dict]:
    """Top-3 land cover transitions."""
    tr = analysis.get("landcover_transition") or {}
    changes: list[dict] = []

    if isinstance(tr.get("top_changes"), list):
        for item in tr["top_changes"]:
            changes.append({
                "from":    item.get("from") or item.get("from_class") or "?",
                "to":      item.get("to")   or item.get("to_class")   or "?",
                "area_ha": float(item.get("area_ha") or item.get("area") or 0),
            })
    elif isinstance(tr.get("pairs"), list):
        for pair in tr["pairs"]:
            for flow in (pair.get("flows") or []):
                if flow.get("changed"):
                    changes.append({
                        "from":    flow.get("from") or "?",
                        "to":      flow.get("to")   or "?",
                        "area_ha": float(flow.get("area") or flow.get("area_ha") or 0),
                    })

    changes.sort(key=lambda x: x["area_ha"], reverse=True)
    return changes[:3]


def extract_vegetation_stats(analysis: dict) -> dict[str, dict]:
    """Returns {INDEX: {mean, min, max, classification, narrative}} from several response shapes.

    classification/narrative are new, optional fields (backend/vegetation_index_registry.py) —
    older cached analysis payloads without them simply omit those keys.
    """
    veg = analysis.get("vegetation") or {}
    raw = veg.get("statistics") or veg.get("stats") or veg.get("indices") or {}
    stats: dict[str, dict] = {}
    for idx, d in raw.items():
        if isinstance(d, dict):
            entry = {
                "mean": d.get("mean"),
                "min":  d.get("min"),
                "max":  d.get("max"),
            }
            if isinstance(d.get("classification"), dict):
                entry["classification"] = d["classification"]
            if isinstance(d.get("narrative"), str):
                entry["narrative"] = d["narrative"]
            stats[idx] = entry
    return stats


def dominant_vegetation_class(veg_entry: dict) -> tuple[str, dict] | None:
    """Largest-area class from a single index's classification block, if present."""
    classes = (veg_entry or {}).get("classification", {}).get("classes")
    if not classes:
        return None
    return max(classes.items(), key=lambda kv: kv[1].get("area", 0))


def extract_carbon_stats(analysis: dict) -> dict:
    """Returns {mean, total, unit, co2e, dataset, model_name, target_pool}."""
    c = analysis.get("carbon") or {}
    raw = c.get("statistics") or c.get("stats") or c.get("summary") or {}
    return {
        "mean":        raw.get("mean"),
        "total":       raw.get("total") or raw.get("total_carbon") or raw.get("total_tons"),
        "unit":        raw.get("unit") or "Mg C/ha",
        "co2e":        raw.get("co2e") or raw.get("carbon_dioxide_equivalent_tons"),
        "dataset":     c.get("dataset"),
        "model_name":  c.get("model_name"),
        "target_pool": c.get("target_pool"),
    }


def extract_map_layers(payload: dict) -> list[str]:
    """Top-3 map layer display names."""
    layers: list[str] = []
    for item in (get_nested(payload, "agent_report.map_layers") or []):
        name = (item.get("name") or item.get("label") or item.get("type") or str(item)) \
               if isinstance(item, dict) else str(item)
        if name and len(layers) < 3:
            layers.append(name)
    return layers


def extract_limitations(payload: dict) -> list[str]:
    """Up to 5 limitation / warning strings."""
    lims: list[str] = []
    for item in (get_nested(payload, "agent_report.limitations") or []):
        if isinstance(item, str) and item.strip() and len(lims) < 5:
            lims.append(item.strip())
    for w in (payload.get("warnings") or []):
        if isinstance(w, str) and w.strip() and len(lims) < 5:
            lims.append(w.strip())
    return lims


def extract_recommendations(payload: dict) -> list[str]:
    """Up to 5 next-action strings."""
    recs: list[str] = []
    for item in (get_nested(payload, "agent_report.next_actions") or []):
        if isinstance(item, str) and item.strip() and len(recs) < 5:
            recs.append(item.strip())
    return recs


def extract_insights(payload: dict) -> list[str]:
    """Up to 6 insight strings from agent_report."""
    items: list[str] = []
    for item in (get_nested(payload, "agent_report.insights") or []):
        if isinstance(item, str) and item.strip() and len(items) < 6:
            items.append(item.strip())
    return items


# ── Business inference ─────────────────────────────────────────────────────────

def infer_business_findings(n: dict) -> list[str]:
    """Business-oriented key findings inferred from analysis data."""
    insights = n["insights"]
    if insights:
        return insights[:6]

    lc_rows  = n["lc_rows"]
    veg      = n["veg_stats"]
    carbon   = n["carbon"]
    tr_rows  = n["tr_rows"]
    findings: list[str] = []

    if lc_rows:
        dom = lc_rows[0]
        findings.append(
            f"Tutupan lahan dominan: {dom['class']} ({_fmt(dom['percent'], 1)}%, "
            f"{_fmt_area(dom['area_ha'])}) — mengindikasikan karakter utama penggunaan lahan "
            "dan menjadi dasar prioritas monitoring."
        )

    if tr_rows:
        top = tr_rows[0]
        findings.append(
            f"Perubahan tutupan lahan terbesar: dari {top['from']} ke {top['to']} "
            f"seluas {_fmt_area(top['area_ha'])} — perlu diprioritaskan untuk verifikasi lapangan."
        )

    ndvi_key = "NDVI" if "NDVI" in veg else (next(iter(veg), None))
    if ndvi_key:
        m = veg[ndvi_key].get("mean")
        if m is not None:
            mv = float(m)
            health = ("baik — mendukung konservasi dan penilaian ESG" if mv >= 0.5
                      else "sedang — perlu dipantau untuk deteksi risiko degradasi" if mv >= 0.3
                      else "rendah — indikasi potensi degradasi atau perubahan tutupan")
            findings.append(f"Kondisi vegetasi ({ndvi_key} mean: {_fmt(mv, 3)}) tergolong {health}.")

    if carbon.get("mean") is not None:
        findings.append(
            f"Estimasi karbon rata-rata {_fmt(carbon['mean'], 1)} {carbon['unit']} — "
            "dapat digunakan sebagai indikasi nilai ekologis untuk carbon accounting atau ESG reporting."
        )

    if carbon.get("total") is not None:
        findings.append(
            f"Total estimasi karbon: {_fmt(carbon['total'], 0)} ton — relevan untuk "
            "pertimbangan program offset karbon atau pelaporan sustainability."
        )

    if not findings:
        findings.append(
            "Analisis geospasial telah dilakukan. Gunakan hasil ini sebagai "
            "screening awal untuk pengambilan keputusan lanjutan."
        )

    return findings[:6]


def infer_business_metrics(n: dict) -> list[tuple[str, str, str]]:
    """Returns list of (Metric, Value, Business Meaning), max 6."""
    rows: list[tuple[str, str, str]] = []
    lc_data  = n["analysis"].get("landcover") or {}
    lc_rows  = n["lc_rows"]
    veg      = n["veg_stats"]
    carbon   = n["carbon"]
    tr_rows  = n["tr_rows"]

    total_ha = lc_data.get("total_area_ha") or lc_data.get("total_area")
    if total_ha:
        rows.append(("Luas AOI", _fmt_area(total_ha), "Skala area yang dianalisis"))

    if lc_rows:
        dom = lc_rows[0]
        rows.append((
            "Tutupan Lahan Dominan",
            f"{dom['class']} ({_fmt(dom['percent'], 1)}%)",
            "Karakter utama penggunaan lahan — baseline monitoring",
        ))

    if tr_rows:
        top = tr_rows[0]
        rows.append((
            "Perubahan Terbesar",
            f"{top['from']} → {top['to']}: {_fmt_area(top['area_ha'])}",
            "Prioritas verifikasi lapangan atau tindakan mitigasi",
        ))

    ndvi_key = "NDVI" if "NDVI" in veg else (next(iter(veg), None))
    if ndvi_key and veg[ndvi_key].get("mean") is not None:
        rows.append((
            f"{ndvi_key} Rata-rata",
            _fmt(veg[ndvi_key]["mean"], 3),
            "Indikator kesehatan vegetasi — ESG dan risiko degradasi",
        ))
        dominant = dominant_vegetation_class(veg[ndvi_key])
        if dominant:
            label, cls = dominant
            rows.append((
                f"Kelas {ndvi_key} Dominan",
                f"{label} ({_fmt(cls.get('percentage'), 1)}%)",
                "Luas kelas terbesar hasil klasifikasi indeks vegetasi",
            ))

    if carbon.get("mean") is not None:
        rows.append((
            "Estimasi Karbon (mean)",
            f"{_fmt(carbon['mean'], 1)} {carbon['unit']}",
            "Kerapatan stok karbon — screening carbon accounting",
        ))

    if carbon.get("total") is not None:
        rows.append((
            "Total Estimasi Karbon",
            f"{_fmt(carbon['total'], 0)} ton",
            "Total stok — relevan untuk offset / sustainability reporting",
        ))

    return rows[:6]


def infer_recommendations(n: dict) -> list[str]:
    """Use provided recs, or fall back to sensible defaults."""
    recs = n["recommendations"]
    if recs:
        return recs[:5]
    return [
        "Lakukan validasi lapangan pada area dengan perubahan tutupan lahan signifikan.",
        "Gunakan hasil ini sebagai screening awal, bukan satu-satunya dasar keputusan investasi.",
        "Prioritaskan monitoring berkala pada zona perubahan atau degradasi.",
        "Bandingkan hasil dengan batas konsesi atau proyek jika tersedia.",
        "Siapkan layer prioritas untuk diskusi manajemen atau ESG reporting.",
    ]


# ── Normalizer ─────────────────────────────────────────────────────────────────

def normalize_report_payload(payload: dict) -> dict:
    """Defensive normalization — always returns a clean dict."""
    analysis     = payload.get("analysis") or {}
    agent_report = payload.get("agent_report") or {}
    period_raw   = payload.get("period") or {}
    bc_raw       = payload.get("business_context") or {}
    meta         = payload.get("metadata") or {}

    lc_rows     = extract_landcover_rows(analysis)
    tr_rows     = extract_transition_rows(analysis)
    veg_stats   = extract_vegetation_stats(analysis)
    carbon      = extract_carbon_stats(analysis)
    limitations = extract_limitations(payload)
    recs        = extract_recommendations(payload)
    insights    = extract_insights(payload)
    map_layers  = extract_map_layers(payload)

    audience_summary = (agent_report.get("audience_summary")
                        or agent_report.get("summary") or "")
    if isinstance(audience_summary, dict):
        audience_summary = audience_summary.get("text") or audience_summary.get("summary") or ""

    year = (period_raw.get("year") or period_raw.get("end_year")
            or datetime.now(UTC).year)
    n: dict = {
        "title":        payload.get("title") or "SaveGeo Executive Summary",
        "aoi_name":     payload.get("aoi_name") or "Area Studi",
        "generated_by": payload.get("generated_by") or "SaveGeo",
        "generated_at": datetime.now(UTC).strftime("%d %B %Y, %H:%M"),
        "period": {
            "year":       year,
            "start_year": period_raw.get("start_year") or year,
            "end_year":   period_raw.get("end_year") or year,
        },
        "business_context": {
            "project_name": bc_raw.get("project_name") or payload.get("aoi_name") or "—",
            "sector":       bc_raw.get("sector") or "—",
            "objective":    bc_raw.get("objective") or "—",
            "audience":     bc_raw.get("audience") or "management",
        },
        "analysis":     analysis,
        "lc_rows":      lc_rows,
        "tr_rows":      tr_rows,
        "veg_stats":    veg_stats,
        "carbon":       carbon,
        "limitations":  limitations,
        "recommendations": recs,
        "insights":     insights,
        "map_layers":   map_layers,
        "audience_summary": audience_summary,
        "datasets":     meta.get("datasets") or [],
        "models":       meta.get("models") or [],
        "parameters":   meta.get("parameters") or {},
    }

    n["findings"]          = infer_business_findings(n)
    n["metrics"]           = infer_business_metrics(n)
    n["final_recs"]        = infer_recommendations(n)
    return n


# ── Section builders ───────────────────────────────────────────────────────────

def _add_compact_header(doc: Document, n: dict) -> None:
    # Title
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(2)
    run = p.add_run(n["title"])
    _set_font(run, 16, bold=True, color=_C_BLUE)

    # Subtitle line
    period = n["period"]
    period_str = (f"{period['start_year']}–{period['end_year']}"
                  if period["start_year"] != period["end_year"]
                  else str(period["year"]))
    _para(doc, f"AOI / Lokasi: {n['aoi_name']}     |     Periode: {period_str}",
          size_pt=9.5, bold=True, color=_C_DARK, space_after=1)
    _para(doc, f"Dibuat: {n['generated_at']}     —     {n['generated_by']}",
          size_pt=8.5, color=_C_GRAY, space_after=4)

    # Thick separator
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after  = Pt(4)
    pPr = p2._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"),   "single")
    bot.set(qn("w:sz"),    "12")
    bot.set(qn("w:space"), "1")
    bot.set(qn("w:color"), _HEX_BLUE)
    pBdr.append(bot)
    pPr.append(pBdr)


def _add_business_summary(doc: Document, n: dict) -> None:
    _heading(doc, "Business Summary")

    summary = (n["audience_summary"] or "").strip()

    if len(summary) > 20:
        # Break into ≤3 paragraphs by sentence
        sentences = [s.strip() for s in summary.replace("\n", " ").split(".") if s.strip()]
        chunk = max(1, (len(sentences) + 2) // 3)
        paras = [". ".join(sentences[i:i+chunk]) + "." for i in range(0, min(len(sentences), chunk*3), chunk)]
        for txt in paras[:3]:
            _para(doc, txt, size_pt=9.5, space_after=3)
    else:
        # Auto-generate from data
        lc, veg, carbon, tr = n["lc_rows"], n["veg_stats"], n["carbon"], n["tr_rows"]
        aoi = n["aoi_name"]

        # Para 1: land cover context
        if lc:
            dom = lc[0]
            txt = (f"Area studi {aoi} didominasi oleh kelas {dom['class']} "
                   f"({_fmt(dom['percent'], 1)}% area), mengindikasikan karakter utama "
                   "penggunaan lahan yang menjadi dasar penilaian risiko dan peluang.")
        else:
            txt = (f"Analisis geospasial telah dilakukan pada area {aoi} "
                   "untuk mendukung pengambilan keputusan berbasis data.")
        _para(doc, txt, size_pt=9.5, space_after=3)

        # Para 2: vegetation + carbon
        parts: list[str] = []
        ndvi_key = "NDVI" if "NDVI" in veg else next(iter(veg), None)
        if ndvi_key and veg[ndvi_key].get("mean") is not None:
            if veg[ndvi_key].get("narrative"):
                parts.append(veg[ndvi_key]["narrative"])
            else:
                mv = float(veg[ndvi_key]["mean"])
                parts.append(
                    f"Nilai {ndvi_key} rata-rata {_fmt(mv, 3)} menunjukkan kondisi vegetasi "
                    f"{'baik' if mv >= 0.5 else 'moderat' if mv >= 0.3 else 'rendah'}, "
                    "relevan untuk penilaian ESG dan risiko degradasi."
                )
        if carbon.get("mean") is not None:
            parts.append(
                f"Estimasi karbon memberikan indikasi nilai ekologis "
                f"(rata-rata {_fmt(carbon['mean'], 1)} {carbon['unit']}) "
                "yang dapat dipertimbangkan untuk offset, konservasi, atau sustainability reporting."
            )
        if parts:
            _para(doc, " ".join(parts), size_pt=9.5, space_after=3)

        # Para 3: change / urgency
        if tr:
            top = tr[0]
            _para(doc,
                  f"Perubahan tutupan lahan teridentifikasi — perubahan terbesar dari "
                  f"{top['from']} ke {top['to']} seluas {_fmt_area(top['area_ha'])}. "
                  "Area perubahan ini memerlukan verifikasi lapangan sebelum keputusan "
                  "operasional atau investasi.",
                  size_pt=9.5, space_after=3)


def _add_strategic_key_findings(doc: Document, n: dict) -> None:
    _heading(doc, "Strategic Key Findings")
    for f in n["findings"]:
        _bullet(doc, f)


def _add_business_metrics_table(doc: Document, n: dict) -> None:
    _heading(doc, "Business-Relevant Metrics")
    metrics = n["metrics"]
    if not metrics:
        _para(doc, "Data tidak tersedia.", size_pt=9, color=_C_GRAY)
        return

    tbl = doc.add_table(rows=1, cols=3)
    tbl.style = "Table Grid"

    # Header
    hdr = tbl.rows[0]
    _fill_row(hdr, ["Metrik", "Nilai", "Makna Bisnis"],
              [9, 9, 9], bold=True, color=_C_WHITE, bg_hex=_HEX_BLUE)

    for i, (metric, value, meaning) in enumerate(metrics):
        row = tbl.add_row()
        bg  = _HEX_LBLUE if i % 2 == 0 else _HEX_WHITE
        _fill_row(row, [metric, value, meaning], [9, 9, 8.5], bg_hex=bg)

    _para(doc, "", space_before=2, space_after=2)


def _add_business_interpretation(doc: Document, n: dict) -> None:
    _heading(doc, "Business Interpretation by Theme")

    lc, veg, carbon, tr = n["lc_rows"], n["veg_stats"], n["carbon"], n["tr_rows"]

    def _inline_heading_para(label: str, text: str,
                              label_color: RGBColor = _C_BLUE) -> None:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after  = Pt(2)
        r1 = p.add_run(label + ": ")
        _set_font(r1, 9.5, bold=True, color=label_color)
        r2 = p.add_run(text)
        _set_font(r2, 9.5)

    if lc:
        dom = lc[0]
        txt = (f"Dominasi {dom['class']} ({_fmt(dom['percent'], 1)}%) mengindikasikan "
               "penggunaan lahan utama sebagai baseline penilaian risiko perubahan, "
               "compliance, dan prioritas monitoring.")
        if len(lc) > 1:
            txt += (f" Kelas {lc[1]['class']} ({_fmt(lc[1]['percent'], 1)}%) menunjukkan "
                    "diversitas tutupan yang relevan untuk perencanaan strategis.")
        _inline_heading_para("Tutupan Lahan", txt)

    ndvi_key = "NDVI" if "NDVI" in veg else next(iter(veg), None)
    if ndvi_key and veg[ndvi_key].get("mean") is not None:
        mv = float(veg[ndvi_key]["mean"])
        interp = ("baik — mendukung konservasi dan penilaian ESG" if mv >= 0.5
                  else "moderat — perlu pemantauan risiko degradasi" if mv >= 0.3
                  else "rendah — perlu perhatian khusus terkait degradasi")
        _inline_heading_para(
            f"Vegetasi ({ndvi_key})",
            f"Nilai rata-rata {_fmt(mv, 3)} menunjukkan kondisi {interp}."
        )

    if carbon.get("mean") is not None or carbon.get("total") is not None:
        parts: list[str] = []
        if carbon.get("mean") is not None:
            parts.append(f"Kerapatan karbon rata-rata {_fmt(carbon['mean'], 1)} {carbon['unit']}")
        if carbon.get("total") is not None:
            parts.append(f"total estimasi {_fmt(carbon['total'], 0)} ton")
        txt = (". ".join(parts) if parts else "Data karbon tersedia") + \
              " — relevan sebagai screening awal untuk carbon accounting dan sustainability reporting."
        if carbon.get("target_pool"):
            txt += (f" Perhatikan target pool ({carbon['target_pool']}) "
                    "agar tidak dibandingkan langsung dengan pool karbon lain.")
        _inline_heading_para("Karbon", txt, label_color=_C_GREEN)

    if tr:
        top = tr[0]
        _inline_heading_para(
            "Analisis Perubahan",
            f"Perubahan dari {top['from']} ke {top['to']} seluas {_fmt_area(top['area_ha'])} "
            "perlu diprioritaskan — dapat berdampak pada risiko operasional, compliance, atau ESG.",
            label_color=_C_RED,
        )

    if not lc and not veg and carbon.get("mean") is None and not tr:
        _para(doc, "Data interpretasi tidak tersedia. Jalankan analisis terlebih dahulu.",
              size_pt=9, color=_C_GRAY)


def _add_decision_considerations(doc: Document, n: dict) -> None:
    _heading(doc, "Decision Considerations")
    lims = n["limitations"] or [
        "Validasi lapangan diperlukan sebelum pengambilan keputusan operasional.",
        "Gunakan hasil ini sebagai screening awal, bukan satu-satunya dasar keputusan investasi.",
        "Prioritaskan monitoring pada area dengan perubahan tutupan lahan signifikan.",
        "Pertimbangkan resolusi dataset dan kondisi awan dalam interpretasi hasil.",
        "Bandingkan dengan data referensi lapangan jika tersedia.",
    ]
    for lim in lims[:5]:
        _bullet(doc, lim)


def _add_recommended_next_actions(doc: Document, n: dict) -> None:
    _heading(doc, "Recommended Next Actions")
    for rec in n["final_recs"]:
        _bullet(doc, rec)


def _add_technical_notes(doc: Document, n: dict) -> None:
    _heading(doc, "Technical Notes", size_pt=8.5, color=_C_GRAY, border_hex="757575",
             space_before=4)

    lines: list[str] = []
    if n["datasets"]:
        lines.append(f"Dataset: {', '.join(str(d) for d in n['datasets'][:4])}")
    if n["models"]:
        lines.append(f"Model: {', '.join(str(m) for m in n['models'][:2])}")
    period = n["period"]
    lines.append(f"Periode analisis: {period['start_year']}–{period['end_year']}")
    carbon = n["carbon"]
    if carbon.get("target_pool"):
        lines.append(f"Target pool karbon: {carbon['target_pool']}")
    if carbon.get("dataset"):
        lines.append(f"Dataset karbon: {carbon['dataset']}")
    if not lines:
        lines.append("Metadata teknis tidak tersedia.")

    for line in lines[:5]:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after  = Pt(1)
        run = p.add_run(line)
        _set_font(run, 8, color=_C_GRAY)

    # Footer
    p_foot = doc.add_paragraph()
    p_foot.paragraph_format.space_before = Pt(6)
    p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_f = p_foot.add_run(
        f"Generated by {n['generated_by']}  ·  {n['generated_at']}"
    )
    _set_font(run_f, 7.5, italic=True, color=_C_GRAY)


# ── Main entry point ───────────────────────────────────────────────────────────

def build_executive_summary_docx(payload: dict) -> BytesIO:
    """
    Build a ~2-page executive summary DOCX from a SaveGeo analysis payload.
    Returns BytesIO ready for Flask send_file().
    Gracefully handles partial or empty payloads.
    """
    n = normalize_report_payload(payload)

    doc = Document()

    # Slim margins to fit ~2 pages
    for sec in doc.sections:
        sec.top_margin    = Cm(1.5)
        sec.bottom_margin = Cm(1.5)
        sec.left_margin   = Cm(1.8)
        sec.right_margin  = Cm(1.8)

    # Default paragraph font
    doc.styles["Normal"].font.name = _FONT
    doc.styles["Normal"].font.size = Pt(10)

    _add_compact_header(doc, n)
    _add_business_summary(doc, n)
    _add_strategic_key_findings(doc, n)
    _add_business_metrics_table(doc, n)
    _add_business_interpretation(doc, n)
    _add_decision_considerations(doc, n)
    _add_recommended_next_actions(doc, n)
    _add_technical_notes(doc, n)

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
